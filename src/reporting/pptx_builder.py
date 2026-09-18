"""python-pptx report deck assembler creating institutional-grade investor presentations."""
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

logger = logging.getLogger(__name__)

# Institutional Color Palette
RGB_NAVY = RGBColor(15, 23, 42)        # Slate 900
RGB_DARK_BLUE = RGBColor(30, 58, 138)  # Blue 900
RGB_PRIMARY = RGBColor(37, 99, 235)    # Blue 600
RGB_TEAL = RGBColor(13, 148, 136)      # Teal 600
RGB_AMBER = RGBColor(217, 119, 6)      # Amber 600
RGB_SLATE_LIGHT = RGBColor(241, 245, 249) # Slate 100
RGB_TEXT_DARK = RGBColor(30, 41, 59)   # Slate 800
RGB_TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
RGB_WHITE = RGBColor(255, 255, 255)


def create_investor_report_presentation(
    company_data: Dict[str, Any],
    periods_data: List[Dict[str, Any]],
    peer_data: Optional[Dict[str, Any]],
    chart_paths: Dict[str, Path],
    output_path: Path,
) -> Path:
    """Assembles a 4-slide widescreen (16:9) PPTX presentation."""
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    ticker = company_data.get("ticker", "UNKNOWN")
    company_name = company_data.get("name", f"{ticker} Corporation")
    sector = company_data.get("sector", "Information Technology")
    exchange = company_data.get("exchange", "NASDAQ")

    # Latest period info
    latest_p = periods_data[-1] if periods_data else {}
    latest_items = latest_p.get("items", {})
    latest_ratios = latest_p.get("computed_ratios", {})

    # ==========================================
    # SLIDE 1: Title Slide (Executive Dark Theme)
    # ==========================================
    slide_1 = prs.slides.add_slide(blank_layout)
    
    # Dark background banner
    bg = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGB_NAVY
    bg.line.fill.background()

    # Blue accent bar
    accent = slide_1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.8), Inches(0.15), Inches(3.8))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGB_PRIMARY
    accent.line.fill.background()

    # Title box
    tx_box = slide_1.shapes.add_textbox(Inches(1.4), Inches(1.8), Inches(11.0), Inches(3.8))
    tf = tx_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "FINANCIAL RESEARCH AUTOMATION (FRA)"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = RGB_PRIMARY
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = f"{company_name} ({ticker})"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGB_WHITE
    p1.space_after = Pt(10)

    p2 = tf.add_paragraph()
    p2.text = f"Quarterly Financial Performance & Peer Benchmarking Report"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGB_SLATE_LIGHT
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    date_str = datetime.now(timezone.utc).strftime("%B %d, %Y")
    p3.text = f"Exchange: {exchange}  |  Sector: {sector}  |  Generated: {date_str}"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGB_TEXT_MUTED

    # ==========================================
    # SLIDE 2: Financial Summary & KPI Scorecards
    # ==========================================
    slide_2 = prs.slides.add_slide(blank_layout)
    _add_slide_header(slide_2, "Executive Financial Summary", f"{company_name} ({ticker}) - Core Performance Indicators")

    # 4 KPI cards across the top
    kpis = [
        ("Latest Revenue", f"${latest_items.get('revenue', 0.0):,.1f} M", "Quarterly Total", RGB_DARK_BLUE),
        ("Net Profit Margin", f"{(latest_ratios.get('net_margin') or 0.0)*100:.1f}%", "Profitability", RGB_TEAL),
        ("Return on Equity (ROE)", f"{(latest_ratios.get('roe') or 0.0)*100:.1f}%", "Shareholder Return", RGB_AMBER),
        ("Current Ratio", f"{(latest_ratios.get('current_ratio') or 0.0):.2f}x", "Liquidity Buffer", RGB_PRIMARY),
    ]

    card_width = Inches(2.7)
    card_height = Inches(1.5)
    start_x = Inches(0.8)
    gap_x = Inches(0.3)
    card_y = Inches(1.8)

    for i, (label, val, sublabel, color) in enumerate(kpis):
        cx = start_x + i * (card_width + gap_x)
        shape = slide_2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGB_SLATE_LIGHT
        shape.line.color.rgb = color
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        p_sub = tf.paragraphs[0]
        p_sub.text = sublabel.upper()
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = RGB_TEXT_MUTED

        p_val = tf.add_paragraph()
        p_val.text = val
        p_val.font.size = Pt(22)
        p_val.font.bold = True
        p_val.font.color.rgb = color

        p_lbl = tf.add_paragraph()
        p_lbl.text = label
        p_lbl.font.size = Pt(11)
        p_lbl.font.color.rgb = RGB_TEXT_DARK

    # Summary Table of all periods
    table_x = Inches(0.8)
    table_y = Inches(3.7)
    table_w = Inches(11.7)
    table_h = Inches(3.0)

    rows = len(periods_data) + 1
    cols = 7
    table_shape = slide_2.shapes.add_table(rows, cols, table_x, table_y, table_w, table_h)
    table = table_shape.table

    headers = ["Period", "Report Date", "Revenue ($M)", "Net Income ($M)", "YoY Growth", "Net Margin", "ROE"]
    for c_idx, h in enumerate(headers):
        cell = table.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGB_DARK_BLUE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = RGB_WHITE

    for r_idx, p in enumerate(periods_data):
        items = p.get("items", {})
        ratios = p.get("computed_ratios", {})
        yoy = ratios.get("yoy_growth")
        nm = ratios.get("net_margin")
        roe = ratios.get("roe")

        row_vals = [
            f"{p.get('period_type')} {p.get('fiscal_year')}",
            str(p.get("report_date")),
            f"{items.get('revenue', 0.0):,.1f}",
            f"{items.get('net_income', 0.0):,.1f}",
            f"{yoy*100:+.1f}%" if yoy is not None else "N/A",
            f"{nm*100:.1f}%" if nm is not None else "N/A",
            f"{roe*100:.1f}%" if roe is not None else "N/A",
        ]
        for c_idx, val in enumerate(row_vals):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = val
            for prg in cell.text_frame.paragraphs:
                prg.font.size = Pt(10)
                prg.font.color.rgb = RGB_TEXT_DARK

    # ==========================================
    # SLIDE 3: Financial Trend Charts
    # ==========================================
    slide_3 = prs.slides.add_slide(blank_layout)
    _add_slide_header(slide_3, "Financial Trajectory & Trends", "Multi-quarter revenue, net income, and margin performance")

    # Add revenue trend chart
    if "revenue" in chart_paths and chart_paths["revenue"].exists():
        slide_3.shapes.add_picture(
            str(chart_paths["revenue"]),
            Inches(0.8),
            Inches(1.8),
            width=Inches(5.7),
        )

    # Add margins chart
    if "margins" in chart_paths and chart_paths["margins"].exists():
        slide_3.shapes.add_picture(
            str(chart_paths["margins"]),
            Inches(6.8),
            Inches(1.8),
            width=Inches(5.7),
        )

    # Commentary box below charts
    notes_box = slide_3.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.9))
    ntf = notes_box.text_frame
    ntf.word_wrap = True
    np1 = ntf.paragraphs[0]
    np1.text = "Key Analytical Insights:"
    np1.font.size = Pt(11)
    np1.font.bold = True
    np1.font.color.rgb = RGB_DARK_BLUE

    np2 = ntf.add_paragraph()
    latest_rev = latest_items.get('revenue', 0.0)
    nm_val = (latest_ratios.get('net_margin') or 0.0) * 100
    np2.text = f"• {ticker} delivered ${latest_rev:,.1f}M in latest quarterly revenue with a robust net margin of {nm_val:.1f}%.\n• Balance sheet shows stable liquidity with consistent operating cash flow generation."
    np2.font.size = Pt(10)
    np2.font.color.rgb = RGB_TEXT_DARK

    # ==========================================
    # SLIDE 4: Peer Comparison Benchmark
    # ==========================================
    slide_4 = prs.slides.add_slide(blank_layout)
    _add_slide_header(slide_4, "Peer Benchmark & Percentile Standing", f"{ticker} competitive position relative to industry peers")

    # Add peer comparison chart on left
    if "peer" in chart_paths and chart_paths["peer"].exists():
        slide_4.shapes.add_picture(
            str(chart_paths["peer"]),
            Inches(0.8),
            Inches(1.8),
            width=Inches(5.8),
        )

    # Peer metrics table on right
    if peer_data and "companies" in peer_data:
        peer_comps = peer_data["companies"]
        p_table_x = Inches(6.9)
        p_table_y = Inches(1.8)
        p_table_w = Inches(5.6)
        p_table_h = Inches(4.5)

        p_rows = len(peer_comps) + 1
        p_cols = 4
        pt_shape = slide_4.shapes.add_table(p_rows, p_cols, p_table_x, p_table_y, p_table_w, p_table_h)
        pt = pt_shape.table

        pt_headers = ["Ticker", "Net Margin", "ROE", "Percentile Rank"]
        for c_idx, h in enumerate(pt_headers):
            cell = pt.cell(0, c_idx)
            cell.text = h
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGB_DARK_BLUE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.bold = True
                p.font.color.rgb = RGB_WHITE

        for r_idx, comp in enumerate(peer_comps):
            c_ticker = comp.get("ticker", "")
            c_nm = comp.get("net_margin")
            c_roe = comp.get("roe")
            pct = comp.get("percentile_rankings", {}).get("net_margin", 50.0)

            c_vals = [
                c_ticker + (" *" if c_ticker.upper() == ticker.upper() else ""),
                f"{c_nm*100:.1f}%" if c_nm is not None else "N/A",
                f"{c_roe*100:.1f}%" if c_roe is not None else "N/A",
                f"{pct:.0f}th pct",
            ]
            for c_idx, val in enumerate(c_vals):
                cell = pt.cell(r_idx + 1, c_idx)
                cell.text = val
                if c_ticker.upper() == ticker.upper():
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(238, 242, 255)  # Light indigo highlight
                for prg in cell.text_frame.paragraphs:
                    prg.font.size = Pt(10)
                    prg.font.bold = (c_ticker.upper() == ticker.upper())
                    prg.font.color.rgb = RGB_PRIMARY if c_ticker.upper() == ticker.upper() else RGB_TEXT_DARK

    # Save presentation
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    logger.info(f"Presentation saved successfully to {output_path}")
    return output_path


def _add_slide_header(slide, title: str, subtitle: str) -> None:
    """Helper to add consistent top header banner."""
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = RGB_DARK_BLUE

    p_sub = tf.add_paragraph()
    p_sub.text = subtitle
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = RGB_TEXT_MUTED
