"""Unit tests for the reporting layer: Matplotlib charts and PPTX presentation assembly."""
from datetime import date
from pathlib import Path
import pytest
from pptx import Presentation

from src.reporting.charts import (
    generate_revenue_trend_chart,
    generate_margins_chart,
    generate_peer_comparison_chart,
)
from src.reporting.pptx_builder import create_investor_report_presentation


@pytest.fixture
def sample_report_data():
    periods = [
        {
            "period_type": "Q1",
            "fiscal_year": 2024,
            "report_date": date(2023, 6, 30),
            "items": {"revenue": 4500.0, "net_income": 700.0},
            "computed_ratios": {"net_margin": 0.155, "roe": 0.08, "yoy_growth": 0.10},
        },
        {
            "period_type": "Q2",
            "fiscal_year": 2024,
            "report_date": date(2023, 9, 30),
            "items": {"revenue": 4700.0, "net_income": 750.0},
            "computed_ratios": {"net_margin": 0.160, "roe": 0.085, "yoy_growth": 0.12},
        },
    ]
    company = {
        "id": 1,
        "ticker": "INFY",
        "name": "Infosys Limited",
        "sector": "Information Technology",
        "exchange": "NYSE",
    }
    peer_data = {
        "companies": [
            {"ticker": "INFY", "net_margin": 0.16, "roe": 0.085, "percentile_rankings": {"net_margin": 75.0}},
            {"ticker": "TCS", "net_margin": 0.18, "roe": 0.12, "percentile_rankings": {"net_margin": 100.0}},
            {"ticker": "WIPRO", "net_margin": 0.13, "roe": 0.06, "percentile_rankings": {"net_margin": 50.0}},
        ]
    }
    return company, periods, peer_data


def test_charts_generation(tmp_path, sample_report_data):
    _, periods, peer_data = sample_report_data

    rev_chart = tmp_path / "rev.png"
    margin_chart = tmp_path / "margin.png"
    peer_chart = tmp_path / "peer.png"

    generate_revenue_trend_chart(periods, rev_chart)
    assert rev_chart.exists()
    assert rev_chart.stat().st_size > 1000

    generate_margins_chart(periods, margin_chart)
    assert margin_chart.exists()
    assert margin_chart.stat().st_size > 1000

    generate_peer_comparison_chart(peer_data["companies"], "INFY", peer_chart)
    assert peer_chart.exists()
    assert peer_chart.stat().st_size > 1000


def test_pptx_deck_generation(tmp_path, sample_report_data):
    company, periods, peer_data = sample_report_data

    # Generate charts
    charts = {
        "revenue": tmp_path / "rev.png",
        "margins": tmp_path / "margin.png",
        "peer": tmp_path / "peer.png",
    }
    generate_revenue_trend_chart(periods, charts["revenue"])
    generate_margins_chart(periods, charts["margins"])
    generate_peer_comparison_chart(peer_data["companies"], "INFY", charts["peer"])

    output_pptx = tmp_path / "test_report.pptx"
    create_investor_report_presentation(
        company_data=company,
        periods_data=periods,
        peer_data=peer_data,
        chart_paths=charts,
        output_path=output_pptx,
    )

    assert output_pptx.exists()
    assert output_pptx.stat().st_size > 5000

    # Load and verify slides
    prs = Presentation(str(output_pptx))
    assert len(prs.slides) == 4

    # Slide 1: Title
    slide1_text = " ".join([shape.text for shape in prs.slides[0].shapes if shape.has_text_frame])
    assert "Infosys Limited" in slide1_text
    assert "INFY" in slide1_text

    # Slide 2: Summary
    slide2_text = " ".join([shape.text for shape in prs.slides[1].shapes if shape.has_text_frame])
    assert "Executive Financial Summary" in slide2_text
